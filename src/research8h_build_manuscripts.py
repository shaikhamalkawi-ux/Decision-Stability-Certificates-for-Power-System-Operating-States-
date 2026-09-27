"""Build editable research-review documents from explicit, versioned JSON content.

Requires python-docx. Rendering and page inspection are separate acceptance steps.
No scientific calculations or experiment outcomes are inferred by this builder.
"""
from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE


def prop(parent, tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    parent.append(node)
    return node


def direction(paragraph, rtl):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = paragraph._p.get_or_add_pPr()
    # Logical start respects both Word and LibreOffice's paragraph direction.
    ppr.find(qn("w:jc")).set(qn("w:val"), "start")
    prop(ppr, "w:bidi", **{"w:val": "1" if rtl else "0"})


def set_run(run, rtl=False, size=None):
    run.font.name = "Arial" if rtl else "Calibri"
    if size:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = prop(rpr, "w:rFonts")
    for name in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + name), "Arial" if rtl else "Calibri")
    prop(rpr, "w:lang", **{"w:val": "ar-AE" if rtl else "en-US", "w:bidi": "ar-AE"})
    prop(rpr, "w:rtl", **{"w:val": "1" if rtl else "0"})
    if rtl and size:
        prop(rpr, "w:szCs", **{"w:val": int(2 * size)})


def add_text(paragraph, text, rtl=False, size=None):
    # Explicit Latin runs prevent long numbers and identifiers reversing in RTL prose.
    pieces = re.split(r"(\[[0-9]+(?:[–,، ]+[0-9]+)*\]|[A-Za-z0-9][A-Za-z0-9 ._*/%:+\-\[\],–]*[A-Za-z0-9%\]]|[A-Za-z0-9])", str(text)) if rtl else [str(text)]
    for value in pieces:
        if not value:
            continue
        is_arabic = bool(re.search(r"[\u0620-\u064a\u066e-\u06d3]", value))
        if rtl and not is_arabic and re.search(r"[A-Za-z0-9]", value):
            # Unicode isolates keep DOI strings, intervals and citation ranges LTR.
            value = "\u2066" + value + "\u2069"
        if not rtl:
            inline_tokens = re.split(r"(?<![A-Za-z0-9_])(A_X|ℓ_X|r_X|a_X|b_X|s_t|T_edge|Φ_H|Φ_0|L_I|U_I|L_T|U_T|d_i|Y_t|Z_t|U_t|U_\(t−1\)|Eref)(?![A-Za-z0-9_])", value)
            for token in inline_tokens:
                if re.fullmatch(r"(?:A_X|ℓ_X|r_X|a_X|b_X|s_t|T_edge|Φ_H|Φ_0|L_I|U_I|L_T|U_T|d_i|Y_t|Z_t|U_t|U_\(t−1\)|Eref)", token):
                    base, script = ("E", "ref") if token == "Eref" else token.split("_", 1)
                    run = paragraph.add_run(base)
                    set_run(run, size=size)
                    run.font.italic = True
                    run = paragraph.add_run(script.strip("()"))
                    set_run(run, size=size)
                    run.font.subscript = True
                elif token:
                    set_run(paragraph.add_run(token), size=size)
            continue
        run = paragraph.add_run(value)
        set_run(run, rtl=rtl and is_arabic, size=size)


def link(paragraph, label, url, rtl=False):
    relation = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = prop(hyperlink, "w:r")
    rpr = prop(run, "w:rPr")
    prop(rpr, "w:color", **{"w:val": "163D64"})
    prop(rpr, "w:u", **{"w:val": "single"})
    prop(rpr, "w:rFonts", **{"w:ascii": "Arial" if rtl else "Calibri", "w:hAnsi": "Arial" if rtl else "Calibri", "w:cs": "Arial"})
    prop(rpr, "w:rtl", **{"w:val": "1" if rtl else "0"})
    prop(run, "w:t").text = label
    paragraph._p.append(hyperlink)


def math_run(value):
    node = OxmlElement("m:r")
    prop(node, "m:t").text = str(value)
    return node


def math_sequence(parent, sequence):
    for item in sequence:
        if isinstance(item, str):
            parent.append(math_run(item))
        elif item[0] in ("sub", "sup"):
            kind, base, script = item
            if not base:
                raise ValueError("Math script must have an explicit base")
            element = prop(parent, "m:sSub" if kind == "sub" else "m:sSup")
            prop(element, "m:e").append(math_run(base))
            prop(element, "m:sub" if kind == "sub" else "m:sup").append(math_run(script))
        elif item[0] == "frac":
            element = prop(parent, "m:f")
            math_sequence(prop(element, "m:num"), item[1])
            math_sequence(prop(element, "m:den"), item[2])
        else:
            raise ValueError(f"Unsupported math node {item!r}")


def math_plain(nodes):
    output = []
    for node in nodes:
        if isinstance(node, str):
            output.append(node)
        elif node[0] in ("sub", "sup"):
            output.append(node[1] + ("_{" if node[0] == "sub" else "^{") + node[2] + "}")
        else:
            output.append("(" + math_plain(node[1]) + ")/(" + math_plain(node[2]) + ")")
    return "".join(output)


def build(source, output, root):
    data = json.loads(source.read_text(encoding="utf-8"))
    rtl = data["language"] == "ar"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.35)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Caption"):
        style = doc.styles[style_name]
        style.font.name = "Arial" if rtl else "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.13 if rtl else 1.1
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].font.size = Pt(12 if rtl else 11)
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].paragraph_format.keep_with_next = True
    doc.styles["Subtitle"].font.size = Pt(12)
    doc.styles["Subtitle"].paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(10.5)
    doc.styles["Caption"].font.bold = False
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Caption"):
        style = doc.styles[style_name]
        rpr = style.element.get_or_add_rPr()
        prop(rpr, "w:szCs", **{"w:val": int(style.font.size.pt * 2)})
        prop(rpr, "w:bCs", **{"w:val": "1" if style.font.bold else "0"})
    # The runtime's base template can supply a title border; keep the review plain.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.core_properties.title = data["title"]
    doc.core_properties.author = data["author"]
    doc.core_properties.subject = data["subtitle"]
    doc.core_properties.comments = data["status"]
    header = section.header.paragraphs[0]
    direction(header, rtl)
    add_text(header, "مراجعة علمية — مسودة بحث" if rtl else "Temporal feasibility | Research draft", rtl, 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for text, style in ((data["title"], "Title"), (data["subtitle"], "Subtitle"), (data["author"], "Normal"), (data["status"], "Normal")):
        paragraph = doc.add_paragraph(style=style)
        direction(paragraph, rtl)
        add_text(paragraph, text, rtl)
    markdown = ["# " + data["title"], data["subtitle"], data["author"], data["status"]]
    for block_index, block in enumerate(data["blocks"]):
        kind = block["t"]
        next_kind = data["blocks"][block_index + 1]["t"] if block_index + 1 < len(data["blocks"]) else None
        if kind in ("h", "p"):
            paragraph = doc.add_paragraph(style="Heading 1" if kind == "h" else "Normal")
            direction(paragraph, rtl)
            add_text(paragraph, block["text"], rtl)
            if next_kind == "eq":
                paragraph.paragraph_format.keep_with_next = True
            markdown.append(("## " if kind == "h" else "") + block["text"])
        elif kind == "eq":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(9)
            paragraph.paragraph_format.keep_with_next = next_kind == "eq"
            math = OxmlElement("m:oMath")
            math_sequence(math, block["nodes"])
            paragraph._p.append(math)
            add_text(paragraph, "     (" + block["id"] + ")")
            markdown.append("Equation (" + block["id"] + "): " + math_plain(block["nodes"]))
        elif kind == "table":
            paragraph = doc.add_paragraph(style="Caption")
            direction(paragraph, rtl)
            paragraph.paragraph_format.keep_with_next = True
            add_text(paragraph, block["caption"], rtl, 10.5)
            table = doc.add_table(rows=1, cols=len(block["headers"]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            if rtl:
                prop(table._tbl.tblPr, "w:bidiVisual")
            borders = prop(table._tbl.tblPr, "w:tblBorders")
            for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
                prop(borders, "w:" + name, **{"w:val": "single", "w:sz": "4", "w:color": "C4C4C4"})
            for i, width in enumerate(block["widths"]):
                table.columns[i].width = Inches(width)
            rows = [block["headers"]] + block["rows"]
            for row_index, values in enumerate(rows):
                row = table.rows[0] if row_index == 0 else table.add_row()
                prop(row._tr.get_or_add_trPr(), "w:cantSplit")
                if row_index == 0:
                    prop(row._tr.get_or_add_trPr(), "w:tblHeader")
                for i, value in enumerate(values):
                    cell = row.cells[i]
                    cell.width = Inches(block["widths"][i])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(4)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_with_next = row_index < len(rows) - 1
                    cell_rtl = rtl and bool(re.search(r"[\u0600-\u06ff]", str(value)))
                    direction(paragraph, cell_rtl)
                    if not cell_rtl and (row_index or i):
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_text(paragraph, str(value), cell_rtl, 10)
                    if row_index == 0:
                        for run in paragraph.runs:
                            run.bold = True
                        prop(cell._tc.get_or_add_tcPr(), "w:shd", **{"w:fill": "F0F0F0"})
            markdown.append(block["caption"] + "\n\n" + "\n".join(["| " + " | ".join(map(str, block["headers"])) + " |", "| " + " | ".join(["---"] * len(block["headers"])) + " |"] + ["| " + " | ".join(map(str, row)) + " |" for row in block["rows"]]))
        elif kind == "figure":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(root / block["path"]), width=Inches(6.25))
            paragraph = doc.add_paragraph(style="Caption")
            direction(paragraph, rtl)
            add_text(paragraph, block["caption"], rtl, 10.5)
            figure_relative = Path(os.path.relpath(root / block["path"], output.parent)).as_posix()
            markdown.append("![" + block["caption"] + "](" + figure_relative + ")")
        elif kind == "link":
            paragraph = doc.add_paragraph()
            direction(paragraph, rtl)
            link(paragraph, block["label"], block["url"], rtl)
            markdown.append("[" + block["label"] + "](" + block["url"] + ")")
        else:
            raise ValueError(f"Unknown block type: {kind}")
    for index, (citation, url) in enumerate(data["references"], 1):
        paragraph = doc.add_paragraph()
        direction(paragraph, False)
        paragraph.paragraph_format.space_after = Pt(8)
        add_text(paragraph, f"[{index}] {citation} ", False, 10.5)
        link(paragraph, url.replace("https://", ""), url)
        markdown.append(f"[{index}] {citation} [{url.replace('https://', '')}]({url})")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    output.with_suffix(".md").write_text("\n\n".join(markdown) + "\n", encoding="utf-8")
    print(json.dumps({"docx": str(output), "blocks": len(data["blocks"]), "references": len(data["references"]), "rtl": rtl}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--root", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output, args.root)
